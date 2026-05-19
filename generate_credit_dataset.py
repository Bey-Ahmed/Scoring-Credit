"""
Génération d'un jeu de données synthétique pour le projet de scoring crédit
ENM860 - Projet Rioux - Été 2026

Contexte: Caisse populaire québécoise, demandes de prêt personnel 5k-50k$.
Structure générative: 4 profils latents d'emprunteurs + corrélations réalistes.
Cible: défaut de paiement à 24 mois, générée par logistique latente.
"""

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Reproductibilité
SEED = 42
rng = np.random.default_rng(SEED)

# ---------------------------------------------------------------------------
# 1. DÉFINITION DES 4 PROFILS LATENTS
# ---------------------------------------------------------------------------
# Chaque profil a une part dans la population et des distributions propres.
# La somme des parts doit faire 1.0.

profils = {
    "JeunePro": {
        "part": 0.27,
        # Numériques: (moyenne, ecart_type) pour les variables continues
        "age": (30, 3.5),
        "revenu": (62000, 12000),
        "anciennete_emploi": (4, 2),
        "anciennete_residence": (3, 2),
        "score_credit_base": (720, 50),
        "nb_demandes": (1.5, 1.2),
        "tdsr_base": (32, 8),           # ratio endettement %
        "montant": (15000, 6000),
        "duree": (36, 12),
        # Catégorielles: distributions de probabilités
        "emploi": {"Permanent": 0.75, "Temporaire": 0.15, "Autonome": 0.08, "SansEmploi": 0.02},
        "education": {"Secondaire": 0.10, "Collegial": 0.30, "Universitaire": 0.60},
        "logement": {"Locataire": 0.65, "Proprio_hypo": 0.15, "Proprio_sans": 0.02, "Chez_parents": 0.18},
        "etat_civil": {"Celibataire": 0.55, "Marie": 0.10, "Conjoint_fait": 0.30, "Divorce": 0.05},
        "objet_pret": {"Auto": 0.25, "Etudes": 0.15, "Consolidation": 0.25, "Voyage": 0.20, "Renovation": 0.15},
        "region": {"Montreal": 0.55, "Quebec": 0.20, "Urbain_autre": 0.20, "Rural": 0.05},
        "client_existant": {"Oui": 0.40, "Non": 0.60},
    },
    "FamilleBanlieue": {
        "part": 0.30,
        "age": (45, 6),
        "revenu": (95000, 18000),
        "anciennete_emploi": (12, 5),
        "anciennete_residence": (9, 5),
        "score_credit_base": (760, 45),
        "nb_demandes": (0.8, 0.9),
        "tdsr_base": (38, 7),
        "montant": (25000, 9000),
        "duree": (48, 14),
        "emploi": {"Permanent": 0.88, "Temporaire": 0.04, "Autonome": 0.07, "SansEmploi": 0.01},
        "education": {"Secondaire": 0.20, "Collegial": 0.35, "Universitaire": 0.45},
        "logement": {"Locataire": 0.10, "Proprio_hypo": 0.70, "Proprio_sans": 0.18, "Chez_parents": 0.02},
        "etat_civil": {"Celibataire": 0.10, "Marie": 0.55, "Conjoint_fait": 0.25, "Divorce": 0.10},
        "objet_pret": {"Auto": 0.30, "Etudes": 0.10, "Consolidation": 0.15, "Voyage": 0.10, "Renovation": 0.35},
        "region": {"Montreal": 0.25, "Quebec": 0.20, "Urbain_autre": 0.45, "Rural": 0.10},
        "client_existant": {"Oui": 0.70, "Non": 0.30},
    },
    "AutonomeInstable": {
        "part": 0.23,
        "age": (38, 7),
        "revenu": (58000, 22000),       # plus variable
        "anciennete_emploi": (3, 2.5),  # courte
        "anciennete_residence": (4, 3),
        "score_credit_base": (640, 70), # plus bas et plus variable
        "nb_demandes": (3.5, 2),        # plus de demandes récentes
        "tdsr_base": (48, 10),          # ratio plus élevé
        "montant": (18000, 8000),
        "duree": (48, 16),
        "emploi": {"Permanent": 0.10, "Temporaire": 0.20, "Autonome": 0.65, "SansEmploi": 0.05},
        "education": {"Secondaire": 0.30, "Collegial": 0.40, "Universitaire": 0.30},
        "logement": {"Locataire": 0.55, "Proprio_hypo": 0.30, "Proprio_sans": 0.05, "Chez_parents": 0.10},
        "etat_civil": {"Celibataire": 0.30, "Marie": 0.25, "Conjoint_fait": 0.25, "Divorce": 0.20},
        "objet_pret": {"Auto": 0.15, "Etudes": 0.05, "Consolidation": 0.55, "Voyage": 0.10, "Renovation": 0.15},
        "region": {"Montreal": 0.35, "Quebec": 0.20, "Urbain_autre": 0.30, "Rural": 0.15},
        "client_existant": {"Oui": 0.30, "Non": 0.70},
    },
    "RetraiteProprio": {
        "part": 0.20,
        "age": (66, 3.5),
        "revenu": (48000, 11000),
        "anciennete_emploi": (0.5, 0.5),  # retraités: peu d'emploi courant
        "anciennete_residence": (20, 8),
        "score_credit_base": (780, 40),
        "nb_demandes": (0.4, 0.6),
        "tdsr_base": (28, 7),
        "montant": (12000, 5000),
        "duree": (36, 12),
        "emploi": {"Permanent": 0.05, "Temporaire": 0.05, "Autonome": 0.05, "SansEmploi": 0.85},
        "education": {"Secondaire": 0.45, "Collegial": 0.30, "Universitaire": 0.25},
        "logement": {"Locataire": 0.15, "Proprio_hypo": 0.10, "Proprio_sans": 0.73, "Chez_parents": 0.02},
        "etat_civil": {"Celibataire": 0.10, "Marie": 0.55, "Conjoint_fait": 0.10, "Divorce": 0.25},
        "objet_pret": {"Auto": 0.20, "Etudes": 0.02, "Consolidation": 0.13, "Voyage": 0.25, "Renovation": 0.40},
        "region": {"Montreal": 0.20, "Quebec": 0.25, "Urbain_autre": 0.35, "Rural": 0.20},
        "client_existant": {"Oui": 0.85, "Non": 0.15},
    },
}

# Vérification: somme des parts
assert abs(sum(p["part"] for p in profils.values()) - 1.0) < 1e-9

N_TOTAL = 1000

# ---------------------------------------------------------------------------
# 2. GÉNÉRATION PROFIL PAR PROFIL
# ---------------------------------------------------------------------------

def tirage_categoriel(distribution, n, rng):
    """Tire n valeurs depuis une distribution de probabilités."""
    cats = list(distribution.keys())
    probs = list(distribution.values())
    return rng.choice(cats, size=n, p=probs)


def generer_profil(nom, params, n, rng):
    """Génère n observations pour un profil donné."""
    # Numériques de base (avec bornes minimales pour réalisme)
    age = np.clip(rng.normal(*params["age"], n), 21, 75).round().astype(int)
    revenu = np.clip(rng.normal(*params["revenu"], n), 18000, 250000).round(-2).astype(int)
    anciennete_emploi = np.clip(rng.normal(*params["anciennete_emploi"], n), 0, 40).round(1)
    anciennete_residence = np.clip(rng.normal(*params["anciennete_residence"], n), 0, 40).round(1)
    nb_demandes = np.clip(rng.normal(*params["nb_demandes"], n), 0, 12).round().astype(int)
    montant = np.clip(rng.normal(*params["montant"], n), 5000, 50000).round(-2).astype(int)
    duree = np.clip(rng.normal(*params["duree"], n), 12, 84).round().astype(int)

    # Corrélations injectées:
    # Score de crédit dépend du score_base mais aussi de variables corrélées
    score_base = rng.normal(*params["score_credit_base"], n)
    # ajustements:
    # +ancienneté emploi -> + score (effet stabilité)
    # +nb demandes -> - score (effet signal de risque)
    # +revenu (centré) -> + score modeste
    score_credit = (
        score_base
        + 1.5 * (anciennete_emploi - anciennete_emploi.mean())
        - 8 * (nb_demandes - nb_demandes.mean())
        + 0.0003 * (revenu - revenu.mean())
    )
    score_credit = np.clip(score_credit, 300, 900).round().astype(int)

    # Ratio d'endettement (TDSR) dépend de:
    # - tdsr_base du profil
    # - revenu (plus haut revenu -> tdsr plus faible, effet inversé)
    # - montant demandé (plus gros montant -> tdsr plus élevé)
    tdsr_base = rng.normal(*params["tdsr_base"], n)
    tdsr = (
        tdsr_base
        - 0.00008 * (revenu - revenu.mean())
        + 0.0002 * (montant - montant.mean())
    )
    tdsr = np.clip(tdsr, 5, 75).round(1)

    # Catégorielles
    emploi = tirage_categoriel(params["emploi"], n, rng)
    education = tirage_categoriel(params["education"], n, rng)
    logement = tirage_categoriel(params["logement"], n, rng)
    etat_civil = tirage_categoriel(params["etat_civil"], n, rng)
    objet_pret = tirage_categoriel(params["objet_pret"], n, rng)
    region = tirage_categoriel(params["region"], n, rng)
    client_existant = tirage_categoriel(params["client_existant"], n, rng)

    df = pd.DataFrame({
        "Profil_Latent": nom,  # gardé pour validation, à retirer pour livraison
        "Age": age,
        "Revenu_Annuel": revenu,
        "Ratio_Endettement_TDSR": tdsr,
        "Anciennete_Emploi_Ans": anciennete_emploi,
        "Anciennete_Residence_Ans": anciennete_residence,
        "Montant_Demande": montant,
        "Duree_Pret_Mois": duree,
        "Score_Credit": score_credit,
        "Nb_Demandes_Recentes": nb_demandes,
        "Statut_Emploi": emploi,
        "Niveau_Education": education,
        "Statut_Residentiel": logement,
        "Etat_Civil": etat_civil,
        "Objet_Pret": objet_pret,
        "Region": region,
        "Client_Existant": client_existant,
    })
    return df


# Répartition exacte des effectifs par profil
effectifs = {nom: int(round(p["part"] * N_TOTAL)) for nom, p in profils.items()}
# Ajustement si l'arrondi ne tombe pas pile sur N_TOTAL
diff = N_TOTAL - sum(effectifs.values())
if diff != 0:
    premier = next(iter(effectifs))
    effectifs[premier] += diff

dfs = [generer_profil(nom, profils[nom], effectifs[nom], rng) for nom in profils]
df = pd.concat(dfs, ignore_index=True)

# Mélange des lignes (sinon les profils sont rangés en blocs)
df = df.sample(frac=1, random_state=SEED).reset_index(drop=True)

# Identifiant client
df.insert(0, "ID_Client", [f"C{str(i+1).zfill(5)}" for i in range(len(df))])


# ---------------------------------------------------------------------------
# 3. GÉNÉRATION DE LA CIBLE: défaut à 24 mois (logistique latente)
# ---------------------------------------------------------------------------
# logit(p) = b0 + b1*TDSR_norm + b2*Score_norm + b3*AncienEmploi_norm
#           + b4*Revenu_norm + b5*NbDemandes_norm + bruit
#
# Signes attendus:
#  + TDSR -> + risque  (b1 > 0)
#  + Score -> - risque (b2 < 0)
#  + Ancien emploi -> - risque (b3 < 0)
#  + Revenu -> - risque (b4 < 0)
#  + Nb demandes -> + risque (b5 > 0)

def standardiser(x):
    return (x - x.mean()) / x.std()

z_tdsr = standardiser(df["Ratio_Endettement_TDSR"])
z_score = standardiser(df["Score_Credit"])
z_emploi = standardiser(df["Anciennete_Emploi_Ans"])
z_revenu = standardiser(df["Revenu_Annuel"])
z_demandes = standardiser(df["Nb_Demandes_Recentes"])

# Coefficients calibrés pour obtenir un taux de défaut ~17%
b0 = -2.6
b1 = 0.75    # TDSR
b2 = -1.05   # Score
b3 = -0.50   # Ancienneté emploi
b4 = -0.30   # Revenu
b5 = 0.55    # Nb demandes

logit = b0 + b1*z_tdsr + b2*z_score + b3*z_emploi + b4*z_revenu + b5*z_demandes
# Bruit gaussien modéré pour éviter une séparabilité parfaite
logit = logit + rng.normal(0, 0.5, size=len(df))

prob_defaut = 1 / (1 + np.exp(-logit))
defaut = (rng.uniform(0, 1, size=len(df)) < prob_defaut).astype(int)

df["Defaut_24mois"] = defaut

# Retrait du profil latent pour le livrable final
df_livraison = df.drop(columns=["Profil_Latent"])

print(f"Taux de défaut global: {defaut.mean():.1%}")
print(f"Effectifs par profil:\n{df['Profil_Latent'].value_counts()}")
print(f"\nTaux de défaut par profil:")
print(df.groupby("Profil_Latent")["Defaut_24mois"].agg(["mean", "count"]))


# ---------------------------------------------------------------------------
# 4. EXPORT EXCEL MULTI-FEUILLES
# ---------------------------------------------------------------------------

output_path = "/mnt/user-data/outputs/scoring_credit_dataset.xlsx"

# Feuille 1: les données
# Feuille 2: dictionnaire de variables
# Feuille 3: méthodologie de génération
# Feuille 4: statistiques descriptives (résumé)

# Dictionnaire de variables
dictionnaire = pd.DataFrame([
    ["ID_Client", "Texte", "Identifiant unique du client (C00001-C01000)", "Identifiant"],
    ["Age", "Numérique (entier)", "Âge du demandeur en années (21-75)", "Stabilité financière"],
    ["Revenu_Annuel", "Numérique (entier)", "Revenu annuel brut déclaré en CAD (18 000-250 000)", "Capacité de remboursement"],
    ["Ratio_Endettement_TDSR", "Numérique (décimal)", "Total Debt Service Ratio en % (5-75)", "Risque de surendettement"],
    ["Anciennete_Emploi_Ans", "Numérique (décimal)", "Années à l'emploi actuel (0-40)", "Stabilité professionnelle"],
    ["Anciennete_Residence_Ans", "Numérique (décimal)", "Années à l'adresse actuelle (0-40)", "Stabilité résidentielle"],
    ["Montant_Demande", "Numérique (entier)", "Montant du prêt demandé en CAD (5 000-50 000)", "Caractéristique du prêt"],
    ["Duree_Pret_Mois", "Numérique (entier)", "Durée du prêt en mois (12-84)", "Caractéristique du prêt"],
    ["Score_Credit", "Numérique (entier)", "Score de crédit Equifax simulé (300-900)", "Historique de crédit"],
    ["Nb_Demandes_Recentes", "Numérique (entier)", "Nombre d'enquêtes de crédit derniers 6 mois (0-12)", "Signal de comportement"],
    ["Statut_Emploi", "Catégorielle", "Permanent / Temporaire / Autonome / SansEmploi", "Type de revenu"],
    ["Niveau_Education", "Catégorielle", "Secondaire / Collegial / Universitaire", "Capital humain"],
    ["Statut_Residentiel", "Catégorielle", "Locataire / Proprio_hypo / Proprio_sans / Chez_parents", "Patrimoine immobilier"],
    ["Etat_Civil", "Catégorielle", "Celibataire / Marie / Conjoint_fait / Divorce", "Situation familiale"],
    ["Objet_Pret", "Catégorielle", "Auto / Etudes / Consolidation / Voyage / Renovation", "But du financement"],
    ["Region", "Catégorielle", "Montreal / Quebec / Urbain_autre / Rural", "Localisation géographique"],
    ["Client_Existant", "Catégorielle", "Oui / Non (déjà membre de la caisse)", "Relation client"],
    ["Defaut_24mois", "Binaire (cible)", "1 = défaut de paiement à 24 mois, 0 = remboursement normal", "Variable à prédire"],
], columns=["Variable", "Type", "Description", "Rôle métier"])

# Méthodologie
methodologie = pd.DataFrame([
    ["Contexte", "Caisse populaire québécoise de taille moyenne souhaitant industrialiser l'évaluation des demandes de prêt personnel (5 000-50 000 $)."],
    ["Objectif", "Système d'aide à la décision pour le scoring de risque de crédit, en remplacement d'une évaluation manuelle hétérogène."],
    ["Taille", f"{N_TOTAL} demandes de prêt simulées."],
    ["Variables", "17 variables explicatives (9 numériques, 7 catégorielles, 1 identifiant) + 1 variable cible binaire."],
    ["Structure générative", "4 profils latents d'emprunteurs représentant des segments réalistes de clientèle: Jeune professionnel urbain (27%), Famille établie banlieue (30%), Travailleur autonome instable (23%), Retraité propriétaire (20%)."],
    ["Corrélations injectées", "Score de crédit corrélé positivement avec ancienneté emploi et revenu, négativement avec nb de demandes récentes. TDSR corrélé négativement avec revenu et positivement avec montant demandé. Les distributions catégorielles sont cohérentes avec chaque profil (ex: famille banlieue majoritairement propriétaire avec hypothèque)."],
    ["Variable cible", "Défaut à 24 mois généré par une fonction logistique latente avec bruit gaussien modéré (σ=0.5). Coefficients calibrés: TDSR (+), Score (-), Ancienneté emploi (-), Revenu (-), Nb demandes (+)."],
    ["Taux de défaut visé", "Autour de 15-20%, conforme au crédit non garanti des institutions financières québécoises."],
    ["Sources d'inspiration", "Structure des variables inspirée du Statlog German Credit Data (UCI ML Repository). Plages de valeurs calibrées sur les rapports publics relatifs à l'endettement des ménages québécois."],
    ["Reproductibilité", f"Seed numpy = {SEED}. Le script de génération est versionné sur GitHub."],
], columns=["Élément", "Description"])

# Stats descriptives
num_cols = ["Age", "Revenu_Annuel", "Ratio_Endettement_TDSR", "Anciennete_Emploi_Ans",
            "Anciennete_Residence_Ans", "Montant_Demande", "Duree_Pret_Mois",
            "Score_Credit", "Nb_Demandes_Recentes"]
stats_num = df_livraison[num_cols].describe().round(2).T.reset_index()
stats_num.columns = ["Variable", "N", "Moyenne", "Écart-type", "Min", "Q1", "Médiane", "Q3", "Max"]

# Stats catégorielles: distributions
cat_cols = ["Statut_Emploi", "Niveau_Education", "Statut_Residentiel", "Etat_Civil",
            "Objet_Pret", "Region", "Client_Existant"]
stats_cat_rows = []
for col in cat_cols:
    counts = df_livraison[col].value_counts(normalize=True).sort_index()
    for modalite, freq in counts.items():
        stats_cat_rows.append([col, modalite, df_livraison[col].value_counts()[modalite], f"{freq:.1%}"])
stats_cat = pd.DataFrame(stats_cat_rows, columns=["Variable", "Modalité", "Effectif", "Fréquence"])

# Taux de défaut par catégorie clé
defaut_par_cat_rows = []
for col in cat_cols:
    grp = df_livraison.groupby(col)["Defaut_24mois"].agg(["mean", "count"])
    for modalite, row in grp.iterrows():
        defaut_par_cat_rows.append([col, modalite, int(row["count"]), f"{row['mean']:.1%}"])
defaut_par_cat = pd.DataFrame(defaut_par_cat_rows, columns=["Variable", "Modalité", "Effectif", "Taux de défaut"])

# Écriture du fichier Excel
with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
    df_livraison.to_excel(writer, sheet_name="Donnees", index=False)
    dictionnaire.to_excel(writer, sheet_name="Dictionnaire_Variables", index=False)
    methodologie.to_excel(writer, sheet_name="Methodologie", index=False)
    stats_num.to_excel(writer, sheet_name="Stats_Numeriques", index=False)
    stats_cat.to_excel(writer, sheet_name="Stats_Categorielles", index=False)
    defaut_par_cat.to_excel(writer, sheet_name="Defaut_par_Categorie", index=False)


# ---------------------------------------------------------------------------
# 5. FORMATAGE ESTHÉTIQUE
# ---------------------------------------------------------------------------

wb = load_workbook(output_path)

header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", start_color="1F4E78")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
cell_font = Font(name="Arial", size=10)
thin_border = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    # En-têtes
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    # Corps
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = cell_font
            cell.border = thin_border
            if isinstance(cell.value, str):
                cell.alignment = Alignment(vertical="center", wrap_text=False)
    # Hauteur de l'en-tête
    ws.row_dimensions[1].height = 28
    # Largeur des colonnes (basée sur le max de la longueur des valeurs, plafonnée)
    for col_idx, col_cells in enumerate(ws.columns, start=1):
        lettre = get_column_letter(col_idx)
        longueurs = [len(str(c.value)) if c.value is not None else 0 for c in col_cells]
        largeur = min(max(longueurs) + 3, 45)
        ws.column_dimensions[lettre].width = max(largeur, 12)
    # Figer la ligne d'en-tête
    ws.freeze_panes = "A2"

# Filtre auto sur la feuille de données
ws_donnees = wb["Donnees"]
ws_donnees.auto_filter.ref = ws_donnees.dimensions

wb.save(output_path)
print(f"\nFichier généré: {output_path}")
print(f"Dimensions du dataset: {df_livraison.shape}")
